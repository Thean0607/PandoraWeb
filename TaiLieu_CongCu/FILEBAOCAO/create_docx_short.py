import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()
doc.add_heading('CHƯƠNG 2: CƠ SỞ LÝ THUYẾT', level=1)

sections = [
    (
        '2.1. C# (C Sharp)',
        'C# là ngôn ngữ lập trình hướng đối tượng mạnh mẽ do Microsoft phát triển, được sử dụng làm ngôn ngữ chủ đạo ở phía máy chủ (Backend). Ngôn ngữ này giúp xử lý các luồng nghiệp vụ phức tạp của dự án như quy trình đăng nhập, giỏ hàng, thanh toán và quản lý đơn hàng một cách nhanh chóng, an toàn và tối ưu.'
    ),
    (
        '2.2. ASP.NET MVC 5',
        'ASP.NET MVC 5 là nền tảng phát triển web phân tách ứng dụng thành ba thành phần: Model (Dữ liệu), View (Giao diện) và Controller (Điều khiển). Kiến trúc này giúp mã nguồn dự án trở nên gọn gàng, dễ bảo trì, và cho phép các thành viên trong nhóm có thể làm việc song song hiệu quả.'
    ),
    (
        '2.3. Microsoft SQL Server',
        'Microsoft SQL Server là hệ quản trị cơ sở dữ liệu quan hệ mạnh mẽ, đảm bảo tính toàn vẹn và an toàn dữ liệu cao. Trong dự án, nó chịu trách nhiệm lưu trữ tập trung toàn bộ thông tin quan trọng như tài khoản người dùng, danh sách sản phẩm và lịch sử giao dịch mua hàng.'
    ),
    (
        '2.4. Entity Framework 6 (EF6)',
        'Entity Framework 6 là công cụ ORM (Object-Relational Mapping) tự động ánh xạ cơ sở dữ liệu thành các đối tượng C#. Dự án sử dụng EF6 theo mô hình Code-First, cho phép thao tác trực tiếp với cơ sở dữ liệu thông qua mã C# mà không cần viết các câu lệnh SQL thủ công.'
    ),
    (
        '2.5. LINQ (Language Integrated Query)',
        'LINQ cung cấp cú pháp truy vấn dữ liệu trực tiếp trong C# một cách ngắn gọn và an toàn. Việc tích hợp LINQ giúp lọc và sắp xếp dữ liệu hiệu quả, đồng thời tự động tham số hóa dữ liệu để ngăn chặn triệt để các lỗ hổng bảo mật như SQL Injection.'
    ),
    (
        '2.6. Cloudinary API',
        'Cloudinary là nền tảng quản lý và lưu trữ tài nguyên đa phương tiện trên đám mây. Dự án tích hợp Cloudinary để tự động upload, tối ưu hóa dung lượng (nén ảnh) và phân phối hình ảnh sản phẩm qua mạng CDN, giúp trang web tải nhanh và mượt mà hơn.'
    ),
    (
        '2.7. HTML5',
        'HTML5 là ngôn ngữ đánh dấu siêu văn bản định hình cấu trúc cốt lõi của website. Nhờ các thẻ ngữ nghĩa (Semantic Tags) chuẩn xác, HTML5 giúp mã nguồn giao diện rõ ràng, thân thiện với trình duyệt và tối ưu hóa tốt cho các công cụ tìm kiếm (SEO).'
    ),
    (
        '2.8. CSS3',
        'CSS3 chịu trách nhiệm định dạng và tạo kiểu dáng thẩm mỹ cho toàn bộ giao diện của dự án. Với CSS3, website được áp dụng nhiều hiệu ứng thiết kế hiện đại như đổ bóng, chuyển động mượt mà và hiệu ứng kính mờ (Glassmorphism), mang lại trải nghiệm thị giác cao cấp.'
    ),
    (
        '2.9. Bootstrap 5',
        'Bootstrap 5 là framework CSS mã nguồn mở giúp xây dựng giao diện tương thích với mọi kích thước màn hình (Responsive Design). Hệ thống lưới linh hoạt và các thành phần giao diện có sẵn của Bootstrap giúp tiết kiệm đáng kể thời gian thiết kế và đảm bảo tính đồng nhất.'
    ),
    (
        '2.10. JavaScript & jQuery',
        'JavaScript và thư viện jQuery được sử dụng để bổ sung tính tương tác động cho website mà không cần tải lại trang. Dự án áp dụng jQuery để xử lý các sự kiện click, tạo hiệu ứng mượt mà và kiểm tra tính hợp lệ của dữ liệu form ngay tại máy khách (Client-side Validation).'
    ),
    (
        '2.11. Newtonsoft.Json',
        'Newtonsoft.Json là thư viện tiêu chuẩn trong .NET dùng để xử lý dữ liệu JSON. Thư viện này đảm nhiệm vai trò tuần tự hóa và giải mã dữ liệu một cách cực kỳ nhanh chóng khi giao tiếp giữa Frontend và Backend thông qua các luồng gọi API ngầm (AJAX).'
    )
]

for title, content in sections:
    # 1. Heading
    doc.add_heading(title, level=2)
    
    # 2. Image placeholder
    p_img = doc.add_paragraph()
    p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_img = p_img.add_run('[HÌNH ẢNH LOGO BẠN CHÈN VÀO ĐÂY (Canh giữa)]')
    run_img.italic = True
    run_img.bold = True
    
    # 3. Short content (2-4 lines)
    p = doc.add_paragraph(content)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        
# Configure style (font Times New Roman, size 13)
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        
output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Chuong2_DaSua_ChuanForm.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
