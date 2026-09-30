import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()
doc.add_heading('CHƯƠNG 4: TRIỂN KHAI HỆ THỐNG VÀ KẾT QUẢ ĐẠT ĐƯỢC', level=1)

text_content = """
4.1. Môi trường triển khai
(Lưu ý: Tại mục 4.1 này, bạn hãy copy 2 Bảng Cấu hình từ file "Cac_Bang_Bieu_Bao_Cao.docx" dán đè vào đây cho đẹp nhé)

4.1.1. Yêu cầu phần cứng
Để quá trình phát triển (development), biên dịch (build) và vận hành thử nghiệm (testing) hệ thống PandoraWeb diễn ra mượt mà, hạn chế tối đa các tình trạng treo máy hay nghẽn cổ chai tài nguyên, môi trường phần cứng được đề xuất và sử dụng với cấu hình như sau:

* Đối với Máy chủ (Server / Môi trường phát triển cục bộ):
- Bộ vi xử lý (CPU): Khuyến nghị sử dụng Intel Core i5 (thế hệ thứ 8 trở lên) hoặc AMD Ryzen 5 tương đương. CPU đa nhân giúp xử lý tốt các luồng biên dịch mã nguồn của Visual Studio và thực thi các câu lệnh truy vấn phức tạp của SQL Server cùng lúc.
- Bộ nhớ trong (RAM): Tối thiểu 8GB (khuyến nghị 16GB). Dung lượng RAM này đảm bảo có thể chạy song song Visual Studio, SQL Server Management Studio và trình duyệt web với hàng chục tab để tra cứu tài liệu mà không bị tràn bộ nhớ.
- Không gian lưu trữ (Storage): Ổ cứng SSD với dung lượng trống tối thiểu 50GB. Ổ SSD đóng vai trò quan trọng trong việc tăng tốc độ đọc/ghi dữ liệu, giúp thời gian khởi động dự án và tốc độ truy xuất cơ sở dữ liệu nhanh hơn gấp nhiều lần so với ổ HDD truyền thống.
- Mạng (Network): Kết nối Internet băng thông rộng, ổn định để tải các thư viện từ NuGet, cập nhật framework và giao tiếp mượt mà với API lưu trữ đám mây Cloudinary.

* Đối với Máy khách (Client / Người dùng cuối):
- Thiết bị: Hỗ trợ đa dạng thiết bị bao gồm Máy tính để bàn (PC), Máy tính xách tay (Laptop), Máy tính bảng (Tablet) và các dòng Điện thoại di động thông minh (Smartphone).
- Cấu hình: Không yêu cầu phần cứng đặc biệt, chỉ cần thiết bị có khả năng kết nối Internet và hiển thị đồ họa cơ bản để tải hình ảnh trang sức.

4.1.2. Môi trường phần mềm và Công cụ phát triển
Nhằm đảm bảo tính đồng bộ, bảo mật và hiệu năng cao nhất cho dự án, quá trình phát triển hệ thống PandoraWeb được thực hiện trên một hệ sinh thái phần mềm mạnh mẽ chủ yếu đến từ Microsoft:

- Công cụ lập trình (IDE) - Microsoft Visual Studio 2022: Đây là môi trường phát triển tích hợp chính thức và toàn diện nhất dành cho các ứng dụng .NET. Visual Studio 2022 cung cấp hệ thống gợi ý mã thông minh (IntelliSense), công cụ gỡ lỗi (Debugger) cực kỳ mạnh mẽ, và giao diện quản lý cấu trúc thư mục dự án (Solution Explorer) trực quan. Nhờ đó, việc viết mã C#, thiết kế giao diện Razor (.cshtml) và quản lý thư viện NuGet trở nên dễ dàng và tối ưu hơn rất nhiều.

- Hệ quản trị cơ sở dữ liệu - Microsoft SQL Server & SSMS: Hệ thống sử dụng Microsoft SQL Server làm nền tảng lưu trữ toàn bộ dữ liệu nghiệp vụ. Đi kèm với đó là công cụ SQL Server Management Studio (SSMS) - giao diện đồ họa giúp ban quản trị thiết kế sơ đồ cơ sở dữ liệu (Database Diagram), trực tiếp truy vấn dữ liệu bằng T-SQL, cũng như thực hiện các tác vụ sao lưu (Backup) và phục hồi (Restore) dữ liệu một cách an toàn.

- Nền tảng công nghệ lõi: Dự án vận hành trên nền tảng .NET Framework phiên bản 4.7.2, kết hợp cùng kiến trúc ASP.NET MVC 5. Nền tảng này cung cấp thư viện lớp cơ sở đồ sộ, đáp ứng mọi nhu cầu xử lý thuật toán phức tạp của một website thương mại điện tử. Ngoài ra, Entity Framework 6 được sử dụng làm công cụ ORM chuyên trách trong việc thao tác giữa mã C# và dữ liệu SQL Server.

- Trình duyệt kiểm thử và Gỡ lỗi: Google Chrome và Microsoft Edge được sử dụng làm môi trường kiểm thử chính ở phía máy khách. Bộ công cụ dành cho nhà phát triển (Chrome Developer Tools) đóng vai trò then chốt trong việc gỡ lỗi mã JavaScript, kiểm tra cấu trúc HTML/CSS, phân tích lưu lượng mạng (Network Panel) và đặc biệt là kiểm tra tính năng hiển thị đáp ứng (Responsive) trên các kích thước màn hình ảo giả lập điện thoại/tablet.

4.2. Kết quả đạt được
4.2.1. Về mặt giao diện (Front-End)
Dự án đã thiết kế thành công một giao diện thương mại điện tử đồng bộ, hiện đại, mang đậm dấu ấn của thương hiệu trang sức cao cấp. Việc ứng dụng framework Bootstrap 5 giúp website tương thích hoàn hảo (Responsive Web Design) trên đa nền tảng. Các khối nội dung, hình ảnh sản phẩm và thanh điều hướng tự động co giãn, sắp xếp hợp lý trên màn hình điện thoại mà không làm vỡ bố cục, mang lại trải nghiệm xem sản phẩm xuyên suốt cho khách hàng.

4.2.2. Về mặt chức năng (Back-End)
- Phân hệ khách hàng: Hệ thống xử lý trơn tru quy trình mua sắm cốt lõi từ bước tìm kiếm sản phẩm, lọc theo chất liệu, thêm vào giỏ hàng, cho đến xác nhận thanh toán. Các tính năng bảo vệ tài khoản cá nhân, xem lịch sử đơn hàng cũng được đảm bảo tính bảo mật.
- Phân hệ quản trị: Xây dựng được trang Admin Dashboard trực quan, cho phép ban quản trị nắm bắt tình hình doanh thu tức thời. Quy trình quản lý sản phẩm, quản lý biến thể (Size, Giá) và cập nhật tiến trình vận chuyển đơn hàng được số hóa và tự động hóa đáng kể.
- Xử lý tài nguyên đám mây: Triển khai tích hợp thành công Cloudinary API, giúp việc tải lên, tối ưu dung lượng và phân phối hình ảnh diễn ra hoàn toàn tự động, giải phóng dung lượng lưu trữ cục bộ cho máy chủ.

4.3. Đánh giá ưu điểm của hệ thống
- Kiến trúc bền vững: Việc áp dụng chuẩn mô hình MVC giúp phân tách rạch ròi giữa Dữ liệu, Giao diện và Luồng điều khiển. Điều này giúp mã nguồn dự án dễ đọc, dễ bảo trì và có tiềm năng mở rộng các tính năng mới trong tương lai.
- Hiệu suất tải trang: Đẩy toàn bộ hình ảnh lên mạng lưới CDN của Cloudinary giúp giảm tải áp lực cho băng thông server nội bộ, tăng tốc độ phản hồi của trang web một cách ấn tượng.
- Trải nghiệm tương tác mượt mà: Áp dụng kỹ thuật jQuery AJAX ở nhiều chức năng (Thêm giỏ hàng, Thay đổi số lượng) giúp website cập nhật dữ liệu trực tiếp mà không cần phải tải lại toàn trang (reload), tạo cảm giác sử dụng mượt mà như một phần mềm thực thụ.
"""

for line in text_content.strip().split('\n'):
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('4.'):
        if len(line.split('.')) == 3: # 4.1. Môi trường...
            doc.add_heading(line, level=2)
        else:
            doc.add_heading(line, level=3)
    else:
        p = doc.add_paragraph(line)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

# Apply font style
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)

output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Chuong4_KhongTrungLap.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
