import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()
doc.add_heading('CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG', level=1)

text_content = """
3.1. Phân tích yêu cầu hệ thống
3.1.1. Khảo sát hiện trạng và mục tiêu dự án
Trong bối cảnh công nghệ thông tin và thương mại điện tử phát triển bùng nổ, hành vi mua sắm của người tiêu dùng đã dịch chuyển mạnh mẽ từ mua sắm truyền thống sang mua sắm trực tuyến. Đối với mặt hàng trang sức, khách hàng ngày càng đòi hỏi trải nghiệm trực tuyến phải chân thực, thông tin sản phẩm minh bạch và quy trình thanh toán an toàn, bảo mật.
Dự án PandoraWeb được xây dựng nhằm giải quyết bài toán đó. Mục tiêu của dự án là xây dựng một nền tảng thương mại điện tử toàn diện, hiện đại, chuyên cung cấp các sản phẩm trang sức cao cấp. Đối với khách hàng, hệ thống mang đến một không gian mua sắm trực quan, dễ dàng tìm kiếm, so sánh và đặt mua sản phẩm. Đối với ban quản trị, hệ thống cung cấp một bộ công cụ mạnh mẽ để quản lý luồng hàng hóa (nhập, xuất, tồn kho), quản lý trạng thái đơn hàng và theo dõi báo cáo doanh thu theo thời gian thực.

3.1.2. Yêu cầu chức năng
Qua quá trình khảo sát và phân tích nghiệp vụ, hệ thống PandoraWeb được thiết kế với hai phân hệ chính, phục vụ cho hai nhóm đối tượng (Actor) khác nhau: Khách hàng và Quản trị viên.

* Phân hệ dành cho Khách hàng (Customer):
- Quản lý tài khoản: Khách hàng có thể đăng ký tài khoản mới, đăng nhập vào hệ thống, sử dụng tính năng quên/khôi phục mật khẩu. Trong phần hồ sơ, người dùng có thể cập nhật thông tin cá nhân (tên, số điện thoại) và quản lý sổ địa chỉ giao hàng để tiện lợi cho các lần mua sau.
- Tìm kiếm và duyệt sản phẩm: Hệ thống cung cấp công cụ tìm kiếm toàn văn bản (Full-text search) và bộ lọc đa chiều. Khách hàng có thể duyệt qua sản phẩm theo Danh mục (Category), Bộ sưu tập (Collection), hoặc Chất liệu (Material). Tại trang chi tiết, khách hàng có thể xem bộ sưu tập hình ảnh, mô tả chi tiết, giá cả và đọc các Đánh giá (Review) từ những người mua trước.
- Quản lý giỏ hàng và Thanh toán: Tính năng giỏ hàng cho phép thêm, sửa số lượng, và xóa sản phẩm linh hoạt. Khi tiến hành thanh toán (Checkout), hệ thống sẽ tổng hợp giá trị đơn hàng, áp dụng các Mã khuyến mãi (Promotion) hợp lệ, tính phí vận chuyển dựa trên địa chỉ và cung cấp các lựa chọn phương thức thanh toán đa dạng.
- Theo dõi lịch sử giao dịch: Khách hàng có thể truy cập lịch sử mua hàng, xem trạng thái hiện tại của đơn (chờ duyệt, đang giao, hoàn thành, đã hủy) và yêu cầu hủy đơn nếu đơn hàng chưa được xử lý.
- Tính năng tương tác: Thêm sản phẩm vào Danh sách yêu thích (Wishlist) để mua sau, theo dõi các bài viết tin tức (Blog) và tra cứu mục Câu hỏi thường gặp (FAQ).

* Phân hệ dành cho Quản trị viên (Admin/Employee):
- Quản trị Danh mục và Thuộc tính: Admin có quyền quản lý toàn bộ hệ thống phân loại sản phẩm, bao gồm việc Thêm/Sửa/Xóa các Danh mục chính, Bộ sưu tập, Chất liệu và Kích thước (Size).
- Quản trị Sản phẩm và Kho hàng: Cho phép Admin tạo sản phẩm mới, định nghĩa các Biến thể sản phẩm (ProductVariant - ví dụ: cùng một chiếc nhẫn nhưng có size 6, size 7, size 8 với giá và số lượng tồn kho khác nhau). Hình ảnh sản phẩm được hỗ trợ tải lên và quản lý tập trung thông qua dịch vụ đám mây Cloudinary.
- Quản lý Đơn hàng: Cung cấp giao diện trực quan giúp Admin theo dõi toàn bộ đơn hàng phát sinh. Admin thực hiện duyệt đơn, cập nhật quy trình vận chuyển và đánh dấu hoàn thành. Chi tiết từng mặt hàng (OrderItem) trong đơn đều được hiển thị minh bạch để đóng gói chính xác.
- Quản lý Khách hàng và Nhân sự: Admin có thể tra cứu thông tin khách hàng, xem lịch sử mua hàng của họ. Đồng thời, hệ thống hỗ trợ quản lý nhân viên nội bộ và phân quyền (Role) linh hoạt (ai được quyền quản lý sản phẩm, ai được quyền duyệt đơn).
- Báo cáo thống kê: Hệ thống cung cấp biểu đồ doanh thu theo thời gian, thống kê top sản phẩm bán chạy và quản lý tồn kho để hỗ trợ chiến lược kinh doanh.
- Quản lý nội dung (CMS): Hỗ trợ tạo các chiến dịch Flash Sale, viết bài Blog tin tức, thay đổi Banner quảng cáo ngoài trang chủ một cách chủ động.

3.1.3. Yêu cầu phi chức năng
- Hiệu năng (Performance): Thời gian phản hồi của mỗi trang không vượt quá 3 giây. Hình ảnh sản phẩm cần được tải nhanh chóng thông qua việc tối ưu và nén tự động trên mạng phân phối nội dung (CDN) của Cloudinary.
- Bảo mật (Security): Mật khẩu người dùng bắt buộc phải được băm (hash) bằng thuật toán an toàn trước khi lưu vào CSDL. Hệ thống phải tích hợp các cơ chế ngăn chặn các cuộc tấn công phổ biến như SQL Injection, Cross-Site Scripting (XSS) và Cross-Site Request Forgery (CSRF).
- Tính khả dụng và Tương thích (Usability & Compatibility): Giao diện người dùng phải tuân thủ triết lý Responsive Web Design (Sử dụng Bootstrap 5), đảm bảo hiển thị hoàn hảo và không bị vỡ bố cục trên mọi thiết bị (Máy tính bàn, Laptop, Máy tính bảng, Điện thoại thông minh). Quy trình thanh toán phải được thiết kế tinh gọn, tối đa 3 bước.

3.2. Mô hình hóa hệ thống
3.2.1. Biểu đồ Use Case tổng quát
[IMG_USECASE_TONGQUAT]
Biểu đồ Use Case tổng quát mô tả bức tranh toàn cảnh về cách các tác nhân (Actor) tương tác với hệ thống. Khách hàng đóng vai trò là người tiêu thụ các dịch vụ thương mại, trong khi Quản trị viên là người vận hành và duy trì hệ thống dữ liệu.

3.2.2. Sơ đồ luồng nghiệp vụ Mua hàng và Thanh toán (Khách hàng)
[IMG_FLOW_KHACHHANG]
Sơ đồ luồng nghiệp vụ mua hàng (Flowchart) mô tả chi tiết quy trình từng bước khách hàng thực hiện từ khi lựa chọn sản phẩm đến khi hoàn tất đơn đặt hàng:
- Bước 1: Khách hàng tìm kiếm và xem chi tiết một sản phẩm trang sức.
- Bước 2: Khách hàng lựa chọn các thuộc tính (Kích thước, Số lượng) và nhấn "Thêm vào giỏ hàng".
- Bước 3: Khách hàng truy cập Giỏ hàng, kiểm tra lại danh sách sản phẩm và tổng tiền tạm tính.
- Bước 4: Chuyển sang màn hình Thanh toán (Checkout). Tại đây, nếu chưa đăng nhập, hệ thống sẽ yêu cầu đăng nhập tài khoản để quản lý đơn hàng tốt hơn.
- Bước 5: Khách hàng điền thông tin địa chỉ giao hàng, áp dụng mã khuyến mãi (nếu có) và lựa chọn phương thức thanh toán.
- Bước 6: Khách hàng xác nhận đặt hàng. Hệ thống sẽ tiến hành kiểm tra số lượng tồn kho của các sản phẩm.
- Bước 7: Nếu tồn kho hợp lệ, hệ thống tạo Đơn hàng (Order), trừ số lượng tồn kho tương ứng và hiển thị thông báo "Đặt hàng thành công". Nếu sản phẩm đã hết hàng, hệ thống cảnh báo lỗi và yêu cầu khách hàng điều chỉnh lại giỏ hàng.

3.2.3. Sơ đồ luồng nghiệp vụ Xử lý đơn hàng (Quản trị viên)
[IMG_FLOW_ADMIN]
Sơ đồ luồng xử lý đơn hàng mô tả quy trình tiếp nhận và vận hành đơn hàng phía sau hệ thống của Ban quản trị:
- Bước 1: Quản trị viên (hoặc Nhân viên kho) đăng nhập vào hệ thống Admin Dashboard.
- Bước 2: Truy cập vào phân hệ Quản lý đơn hàng, hệ thống hiển thị danh sách các đơn hàng mới nhận (Trạng thái mặc định: Pending).
- Bước 3: Admin chọn một đơn hàng để xem chi tiết bao gồm thông tin khách nhận, danh sách mặt hàng, kích thước và tổng tiền.
- Bước 4: Admin tiến hành kiểm tra hàng hóa thực tế và xác nhận đơn hàng. Trạng thái đơn hàng được chuyển sang "Processing" (Đang xử lý / Đóng gói).
- Bước 5: Sau khi đóng gói hoàn tất và bàn giao cho đơn vị vận chuyển, Admin cập nhật trạng thái đơn thành "Shipped" (Đang giao hàng).
- Bước 6: Khi có xác nhận khách hàng nhận được hàng thành công, Admin cập nhật trạng thái cuối cùng là "Delivered" (Hoàn thành), hệ thống ghi nhận doanh thu. Trường hợp khách từ chối nhận hàng hoặc hết hàng đột xuất, Admin chuyển trạng thái sang "Cancelled" (Đã hủy) và hệ thống tự động hoàn lại số lượng tồn kho.

3.3. Thiết kế Cơ sở dữ liệu (Database Design)
3.3.1. Sơ đồ quan hệ thực thể (ERD)
[IMG_ERD]
Cơ sở dữ liệu của PandoraWeb được thiết kế trên mô hình chuẩn hóa cao (Normal Forms), nhằm giảm thiểu sự dư thừa dữ liệu và đảm bảo tính toàn vẹn tham chiếu thông qua các ràng buộc Khóa chính (Primary Key) và Khóa ngoại (Foreign Key). Sơ đồ ERD minh họa mối liên kết chặt chẽ giữa Hóa đơn - Chi tiết hóa đơn - Biến thể sản phẩm - Khách hàng.

3.3.2. Từ điển dữ liệu (Mô tả các thực thể chính)
- Thực thể Product (Sản phẩm): Chứa thông tin gốc của sản phẩm bao gồm ProductId (Khóa chính), Name (Tên), Description (Mô tả), BasePrice (Giá gốc), CategoryId (Khóa ngoại trỏ đến danh mục), CollectionId (Bộ sưu tập).
- Thực thể ProductVariant (Biến thể sản phẩm): Quản lý các thuộc tính phân nhánh của một sản phẩm. Chứa VariantId, ProductId (Khóa ngoại), SizeId, MaterialId, Price (Giá bán của biến thể này), StockQuantity (Số lượng tồn kho thực tế).
- Thực thể Customer (Khách hàng): Lưu trữ CustomerId, FullName, Email (Unique), PasswordHash (Đã mã hóa), Phone, Status (Hoạt động/Bị khóa).
- Thực thể Order (Đơn hàng gốc): Chứa OrderId, CustomerId (Ai đặt), OrderDate (Ngày đặt), TotalAmount (Tổng tiền), Status (Trạng thái đơn: Pending, Processing, Shipped, Delivered, Cancelled), ShippingAddress (Địa chỉ nhận hàng).
- Thực thể OrderItem (Chi tiết đơn hàng): Lưu trữ các món hàng trong một Order. Chứa OrderItemId, OrderId (Khóa ngoại), VariantId (Sản phẩm biến thể nào), Quantity (Số lượng mua), UnitPrice (Giá tại thời điểm mua).
- Các thực thể Danh mục (Category, Collection, Material, Size): Đóng vai trò là dữ liệu từ điển, hỗ trợ phân cấp, lọc và nhóm sản phẩm.

3.4. Thiết kế giao diện (UI Design)
Hệ thống giao diện được thiết kế tuân theo nguyên tắc UI/UX hiện đại, kết hợp hài hòa giữa màu sắc thương hiệu và khoảng trắng để làm nổi bật sản phẩm.

3.4.1. Giao diện Trang chủ (Home Page)
[IMG_UI_HOME]
Trang chủ là bộ mặt của website. Phần trên cùng (Header) chứa Logo, thanh Tìm kiếm trung tâm, Giỏ hàng và Nút đăng nhập. Ngay bên dưới là hệ thống Banner trình chiếu (Carousel) quảng bá các chiến dịch lớn. Thân trang hiển thị các dải sản phẩm theo tab: Sản phẩm mới (New Arrivals), Sản phẩm bán chạy (Best Sellers) giúp thu hút sự chú ý ngay lập tức.

3.4.2. Giao diện Danh sách sản phẩm (Shop/Category)
[IMG_UI_SHOP]
Giao diện này áp dụng cấu trúc 2 cột. Cột bên trái (Sidebar) tích hợp công cụ lọc mạnh mẽ (Lọc theo khoảng giá, theo Size, Chất liệu). Cột bên phải chiếm diện tích lớn, hiển thị sản phẩm dưới dạng lưới (Grid). Mỗi thẻ sản phẩm (Product Card) hiển thị ảnh sắc nét, tên, giá bán và hiệu ứng trượt chuột (Hover) để hiện nút Thêm vào giỏ (Add to Cart). Hệ thống có tích hợp phân trang (Pagination) ở dưới cùng.

3.4.3. Giao diện Tìm kiếm sản phẩm (Search Results)
[IMG_UI_SEARCH]
Giao diện này xuất hiện khi khách hàng nhập từ khóa vào thanh tìm kiếm. Kết quả trả về được trình bày dưới dạng lưới (Grid), đi kèm với số lượng kết quả tìm thấy và bộ lọc để người dùng nhanh chóng khoanh vùng được sản phẩm mong muốn.

3.4.4. Giao diện Chi tiết sản phẩm (Product Details)
[IMG_UI_DETAIL]
Được thiết kế tối ưu để thúc đẩy chuyển đổi mua hàng. Bên trái là thư viện ảnh sản phẩm (Gallery) cho phép phóng to (Zoom). Bên phải là thông tin chi tiết, giá tiền và bộ chọn Biến thể (Chọn Kích thước, Màu sắc). Nút "Thêm vào giỏ hàng" và "Mua ngay" được thiết kế to, nổi bật. Cuối trang là khu vực Mô tả dài, Đánh giá của khách hàng (Reviews) và gợi ý Sản phẩm liên quan (Related Products).

3.4.5. Giao diện Giỏ hàng (Shopping Cart)
[IMG_UI_CART]
Giao diện giỏ hàng hiển thị danh sách các sản phẩm khách hàng đã chọn mua, kèm theo hình ảnh thu nhỏ, tên sản phẩm, kích thước, đơn giá và tổng tiền. Khách hàng có thể dễ dàng tăng giảm số lượng hoặc xóa sản phẩm khỏi giỏ trực tiếp trên giao diện này trước khi chuyển sang bước thanh toán.

3.4.6. Giao diện Thanh toán (Checkout)
[IMG_UI_CHECKOUT]
Đây là trang then chốt trong quy trình mua hàng. Giao diện được thiết kế tối giản, loại bỏ các thành phần điều hướng rườm rà để khách hàng tập trung điền thông tin giao hàng (Họ tên, Số điện thoại, Địa chỉ). Màn hình hiển thị tóm tắt đơn hàng, form nhập mã giảm giá và cung cấp các phương thức thanh toán một cách minh bạch.

3.4.7. Giao diện Quản lý hồ sơ cá nhân (User Profile & Security)
[IMG_UI_PROFILE]
Khu vực dành riêng cho khách hàng đã đăng nhập, bao gồm các tab chức năng:
- Hồ sơ của tôi (Profile): Xem và chỉnh sửa thông tin cá nhân cơ bản (Họ tên, Email, Số điện thoại).
- Đổi mật khẩu (Change Password): Cung cấp form bảo mật yêu cầu nhập mật khẩu cũ và mật khẩu mới để đảm bảo an toàn tài khoản.
- Lịch sử đơn hàng: Liệt kê danh sách các đơn hàng đã mua và trạng thái xử lý hiện tại của chúng.

3.4.8. Giao diện Quản trị viên (Admin Dashboard)
[IMG_UI_ADMIN]
Giao diện Admin tập trung vào tính hiệu quả và xử lý số liệu. Thanh điều hướng tĩnh nằm bên trái (Sidebar) chứa các menu chức năng: Dashboard, Orders, Products, Customers. Vùng nội dung chính hiển thị các thẻ tóm tắt (Tổng doanh thu, Đơn hàng mới) và các biểu đồ thống kê trực quan. Các bảng dữ liệu sử dụng thư viện DataTables hỗ trợ tìm kiếm, sắp xếp và xuất file Excel nhanh chóng.
"""

for line in text_content.strip().split('\n'):
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('3.'):
        if len(line.split('.')) == 3: # 3.1. Phân tích...
            doc.add_heading(line, level=2)
        else:
            doc.add_heading(line, level=3)
    elif line.startswith('[IMG'):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        text_ph = '[HÌNH ẢNH BẠN CHÈN VÀO ĐÂY (Canh giữa)]'
        if 'USECASE_TONGQUAT' in line: text_ph = '[HÌNH ẢNH BIỂU ĐỒ USE CASE TỔNG QUÁT BẠN CHÈN VÀO ĐÂY]'
        elif 'FLOW_KHACHHANG' in line: text_ph = '[HÌNH ẢNH SƠ ĐỒ LUỒNG (FLOWCHART) MUA HÀNG CỦA KHÁCH BẠN CHÈN VÀO ĐÂY]'
        elif 'FLOW_ADMIN' in line: text_ph = '[HÌNH ẢNH SƠ ĐỒ LUỒNG (FLOWCHART) XỬ LÝ ĐƠN HÀNG BẠN CHÈN VÀO ĐÂY]'
        elif 'ERD' in line: text_ph = '[HÌNH ẢNH BIỂU ĐỒ ERD CƠ SỞ DỮ LIỆU BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_HOME' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN TRANG CHỦ BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_SHOP' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN TRANG DANH SÁCH SẢN PHẨM BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_SEARCH' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN TÌM KIẾM SẢN PHẨM BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_CART' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN GIỎ HÀNG BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_CHECKOUT' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN THANH TOÁN (CHECKOUT) BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_PROFILE' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN QUẢN LÝ HỒ SƠ / ĐỔI MẬT KHẨU BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_DETAIL' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN TRANG CHI TIẾT SẢN PHẨM BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_ADMIN' in line: text_ph = '[HÌNH ẢNH GIAO DIỆN TRANG ADMIN DASHBOARD BẠN CHÈN VÀO ĐÂY]'
        
        run_img = p_img.add_run(text_ph)
        run_img.italic = True
        run_img.bold = True
    else:
        p = doc.add_paragraph(line)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

# Apply font style
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)

output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Chuong3_BanHoanChinh.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
