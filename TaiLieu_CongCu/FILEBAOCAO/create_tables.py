import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Helper function for adding headings
def add_heading(text, level=2):
    p = doc.add_heading(text, level=level)
    return p

# 1. BẢNG MÔI TRƯỜNG TRIỂN KHAI (DÙNG CHO CHƯƠNG 4)
doc.add_heading('1. BẢNG MÔI TRƯỜNG TRIỂN KHAI (Thay thế văn bản mục 4.1)', level=1)

doc.add_paragraph('Bảng 4.1: Cấu hình phần cứng tối thiểu và đề nghị')
table_hw = doc.add_table(rows=1, cols=3)
table_hw.style = 'Table Grid'
hdr = table_hw.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = 'Thiết bị / Linh kiện', 'Yêu cầu tối thiểu (Khách hàng)', 'Yêu cầu đề nghị (Máy chủ / Máy Dev)'

data_hw = [
    ('CPU', 'Không yêu cầu đặc biệt', 'Intel Core i5 / AMD Ryzen 5 trở lên'),
    ('RAM', 'Tối thiểu 2GB', 'Từ 8GB đến 16GB'),
    ('Ổ cứng (Storage)', 'Trống 500MB', 'SSD trống tối thiểu 50GB'),
    ('Kết nối mạng', '3G/4G/Wifi cơ bản', 'Băng thông rộng, cáp quang ổn định'),
    ('Màn hình', 'Smartphone 4 inch trở lên', 'Màn hình Full HD (1920x1080)')
]
for item in data_hw:
    row = table_hw.add_row().cells
    row[0].text, row[1].text, row[2].text = item

doc.add_paragraph('\nBảng 4.2: Môi trường phần mềm và công nghệ sử dụng')
table_sw = doc.add_table(rows=1, cols=2)
table_sw.style = 'Table Grid'
hdr_sw = table_sw.rows[0].cells
hdr_sw[0].text, hdr_sw[1].text = 'Phân loại', 'Công cụ / Công nghệ sử dụng'

data_sw = [
    ('Công cụ lập trình (IDE)', 'Microsoft Visual Studio 2022'),
    ('Hệ quản trị Cơ sở dữ liệu', 'Microsoft SQL Server Management Studio (SSMS)'),
    ('Nền tảng Back-end', '.NET Framework 4.7.2, ASP.NET MVC 5, C#'),
    ('Công nghệ Front-end', 'HTML5, CSS3, Bootstrap 5, jQuery, AJAX'),
    ('Giao tiếp Dữ liệu (ORM)', 'Entity Framework 6 (Code-First), LINQ'),
    ('Lưu trữ đám mây', 'Cloudinary API (Lưu trữ và phân phối hình ảnh)')
]
for item in data_sw:
    row = table_sw.add_row().cells
    row[0].text, row[1].text = item

doc.add_page_break()

# 2. BẢNG TỪ ĐIỂN DỮ LIỆU (DÙNG CHO CHƯƠNG 3)
doc.add_heading('2. CÁC BẢNG TỪ ĐIỂN DỮ LIỆU CSDL (Thay thế văn bản mục 3.3.2)', level=1)
doc.add_paragraph('Dưới đây là đặc tả chi tiết cấu trúc các bảng chính trong Cơ sở dữ liệu của hệ thống PandoraWeb.')

def create_db_table(title, columns_data):
    doc.add_paragraph(title)
    tb = doc.add_table(rows=1, cols=4)
    tb.style = 'Table Grid'
    hdr = tb.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'Tên trường (Column)', 'Kiểu dữ liệu', 'Khóa / Ràng buộc', 'Mô tả ý nghĩa'
    
    for row_data in columns_data:
        row = tb.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = row_data
    doc.add_paragraph('\n')

create_db_table('Bảng 3.1: Đặc tả bảng Khách hàng (Customer)', [
    ('CustomerId', 'int', 'PK, Tự tăng', 'Mã khách hàng duy nhất'),
    ('FullName', 'nvarchar(100)', 'Not Null', 'Họ và tên đầy đủ của khách'),
    ('Email', 'varchar(150)', 'Unique, Not Null', 'Tài khoản đăng nhập / Liên hệ'),
    ('PasswordHash', 'nvarchar(max)', 'Not Null', 'Mật khẩu đã băm bảo mật'),
    ('Phone', 'varchar(20)', 'Null', 'Số điện thoại liên lạc'),
    ('Address', 'nvarchar(255)', 'Null', 'Địa chỉ giao hàng mặc định'),
    ('CreatedAt', 'datetime', 'Default(GETDATE)', 'Ngày tạo tài khoản')
])

create_db_table('Bảng 3.2: Đặc tả bảng Sản phẩm (Product)', [
    ('ProductId', 'int', 'PK, Tự tăng', 'Mã sản phẩm duy nhất'),
    ('Name', 'nvarchar(255)', 'Not Null', 'Tên trang sức'),
    ('Description', 'nvarchar(max)', 'Null', 'Mô tả chi tiết sản phẩm'),
    ('BasePrice', 'decimal(18,2)', 'Not Null', 'Giá niêm yết cơ bản'),
    ('CategoryId', 'int', 'FK (Category)', 'Mã danh mục chứa sản phẩm'),
    ('CollectionId', 'int', 'FK (Collection)', 'Mã bộ sưu tập (nếu có)'),
    ('IsActive', 'bit', 'Default(1)', 'Trạng thái hiển thị web')
])

create_db_table('Bảng 3.3: Đặc tả bảng Biến thể sản phẩm (ProductVariant)', [
    ('VariantId', 'int', 'PK, Tự tăng', 'Mã phân loại sản phẩm (ví dụ nhẫn size 6)'),
    ('ProductId', 'int', 'FK (Product)', 'Mã sản phẩm gốc'),
    ('SizeId', 'int', 'FK (Size)', 'Kích thước của biến thể'),
    ('MaterialId', 'int', 'FK (Material)', 'Chất liệu (Bạc, Vàng...)'),
    ('Price', 'decimal(18,2)', 'Not Null', 'Giá bán cụ thể của biến thể này'),
    ('StockQuantity', 'int', 'Not Null, >= 0', 'Số lượng còn trong kho thực tế')
])

create_db_table('Bảng 3.4: Đặc tả bảng Hóa đơn đặt hàng (Order)', [
    ('OrderId', 'int', 'PK, Tự tăng', 'Mã hóa đơn duy nhất'),
    ('CustomerId', 'int', 'FK (Customer)', 'Mã khách đặt hàng'),
    ('OrderDate', 'datetime', 'Default(GETDATE)', 'Thời điểm đặt hàng'),
    ('TotalAmount', 'decimal(18,2)', 'Not Null', 'Tổng giá trị hóa đơn (đã cộng phí)'),
    ('OrderStatus', 'varchar(50)', 'Not Null', 'Trạng thái (Pending, Shipped, Cancelled...)'),
    ('ShippingAddress', 'nvarchar(500)', 'Not Null', 'Địa chỉ nhận hàng cụ thể'),
    ('PaymentMethod', 'varchar(50)', 'Not Null', 'Phương thức thanh toán (COD, VNPay...)')
])

create_db_table('Bảng 3.5: Đặc tả bảng Chi tiết hóa đơn (OrderItem)', [
    ('OrderItemId', 'int', 'PK, Tự tăng', 'Mã dòng chi tiết'),
    ('OrderId', 'int', 'FK (Order)', 'Mã hóa đơn tổng'),
    ('VariantId', 'int', 'FK (ProductVariant)', 'Sản phẩm/Kích thước cụ thể được mua'),
    ('Quantity', 'int', 'Not Null, > 0', 'Số lượng mua'),
    ('UnitPrice', 'decimal(18,2)', 'Not Null', 'Giá bán chốt tại thời điểm mua hàng')
])

# Apply font style for the whole doc
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)

output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Cac_Bang_Bieu_Bao_Cao.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
