import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()
doc.add_heading('CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG', level=1)

text_content = """
3.1. Phân tích yêu cầu hệ thống
3.1.1. Khảo sát hiện trạng và mục tiêu dự án
Dự án PandoraWeb được xây dựng với mục tiêu cung cấp một nền tảng thương mại điện tử chuyên nghiệp, chuyên phân phối các sản phẩm trang sức. Hệ thống được thiết kế để giải quyết bài toán mua sắm trực tuyến, giúp khách hàng dễ dàng tìm kiếm, xem chi tiết và đặt mua trang sức một cách an toàn, đồng thời giúp chủ cửa hàng quản lý kho hàng, đơn hàng và doanh thu hiệu quả.

3.1.2. Yêu cầu chức năng
Hệ thống được chia thành hai nhóm người dùng chính với các chức năng tương ứng:
* Đối với Khách hàng (Customer):
- Quản lý tài khoản: Đăng ký, đăng nhập, cập nhật thông tin cá nhân và sổ địa chỉ.
- Tra cứu sản phẩm: Xem danh sách sản phẩm theo danh mục (Category), bộ sưu tập (Collection), chất liệu (Material) và xem đánh giá (Review).
- Quản lý giỏ hàng: Thêm, sửa, xóa sản phẩm trong giỏ hàng (Cart).
- Thanh toán: Tiến hành đặt hàng (Order), áp dụng khuyến mãi (Promotion) và theo dõi trạng thái đơn hàng.
- Tính năng khác: Thêm sản phẩm yêu thích (Wishlist), đọc Blog.

* Đối với Quản trị viên (Admin / Employee):
- Quản lý sản phẩm: Thêm, sửa, xóa sản phẩm, biến thể (ProductVariant) và hình ảnh (ProductImage).
- Quản lý danh mục: Quản lý danh mục, bộ sưu tập, chất liệu và kích thước.
- Quản lý đơn hàng: Duyệt đơn, cập nhật trạng thái đơn và xem chi tiết đơn hàng (OrderItem).
- Quản lý người dùng: Quản lý khách hàng, nhân viên và phân quyền (Role).
- Thống kê: Xem báo cáo doanh thu, sản phẩm bán chạy.
- Quản lý nội dung: Quản lý bài viết (Blog), Banner, trang nội dung (Faq, Page).

3.1.3. Yêu cầu phi chức năng
- Hiệu năng (Performance): Tốc độ phản hồi trang nhanh, hình ảnh được tối ưu hóa qua Cloudinary.
- Bảo mật (Security): Mật khẩu mã hóa, chống tấn công SQL Injection và XSS.
- Giao diện (UI/UX): Thiết kế thích ứng (Responsive) trên PC và Mobile, dễ sử dụng, phong cách hiện đại.

3.2. Biểu đồ Use Case (Use Case Diagram)
3.2.1. Biểu đồ Use Case tổng quát
[IMG_USECASE_TONGQUAT]
Biểu đồ Use Case tổng quát mô tả bức tranh toàn cảnh về các tác nhân (Actor) và các nhóm chức năng chính mà hệ thống PandoraWeb cung cấp.

3.2.2. Biểu đồ Use Case Khách hàng
[IMG_USECASE_KHACHHANG]
Khách hàng có thể thực hiện các thao tác từ xem sản phẩm, quản lý giỏ hàng, đến thanh toán và theo dõi đơn hàng một cách xuyên suốt.

3.2.3. Biểu đồ Use Case Quản trị viên (Admin)
[IMG_USECASE_ADMIN]
Quản trị viên có toàn quyền kiểm soát dữ liệu hệ thống, từ việc vận hành luồng sản phẩm, xử lý đơn hàng đến việc theo dõi doanh thu.

3.3. Thiết kế cơ sở dữ liệu (Database Design)
3.3.1. Biểu đồ quan hệ thực thể (ERD)
[IMG_ERD]
Hệ thống cơ sở dữ liệu của PandoraWeb được thiết kế chặt chẽ. Các bảng được liên kết với nhau thông qua hệ thống khóa chính và khóa ngoại để đảm bảo tính toàn vẹn dữ liệu.

3.3.2. Các thực thể (Bảng) chính trong hệ thống
- Bảng Product & ProductVariant: Lưu trữ thông tin sản phẩm và các biến thể (màu sắc, kích thước, giá bán, tồn kho).
- Bảng Customer & Employee: Lưu trữ thông tin tài khoản, thông tin cá nhân và phân quyền truy cập.
- Bảng Order & OrderItem: Lưu trữ thông tin hóa đơn và chi tiết từng mặt hàng được mua trong hóa đơn.
- Bảng Category, Collection, Material, Size: Đóng vai trò là các bảng từ điển dữ liệu, giúp phân loại sản phẩm.
- Bảng Cart & CartItem: Lưu trữ trạng thái giỏ hàng hiện tại của khách hàng.

3.4. Thiết kế giao diện (UI Design)
3.4.1. Giao diện trang chủ (Home Page)
[IMG_UI_HOME]
Trang chủ hiển thị các Banner nổi bật, danh sách danh mục và các sản phẩm bán chạy nhất, mang lại cái nhìn tổng quan.

3.4.2. Giao diện danh sách sản phẩm (Shop)
[IMG_UI_SHOP]
Giao diện hiển thị dạng lưới (Grid), tích hợp bộ lọc bên trái giúp tìm kiếm sản phẩm theo chất liệu, kích thước.

3.4.3. Giao diện chi tiết sản phẩm (Product Details)
[IMG_UI_DETAIL]
Cung cấp hình ảnh trực quan, thông tin chi tiết, và cho phép khách hàng chọn kích thước trước khi thêm vào giỏ.

3.4.4. Giao diện Quản trị viên (Admin Dashboard)
[IMG_UI_ADMIN]
Khu vực dành riêng cho Ban quản trị với các biểu đồ thống kê trực quan và thanh menu điều hướng.
"""

for line in text_content.strip().split('\n'):
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('3.'):
        if len(line.split('.')) == 3: # 3.1. Phân tích...
            doc.add_heading(line, level=2)
        else:
            doc.add_heading(line, level=3)
    elif line.startswith('[IMG'):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        text_ph = '[HÌNH ẢNH BẠN CHÈN VÀO ĐÂY (Canh giữa)]'
        if 'USECASE_TONGQUAT' in line: text_ph = '[HÌNH ẢNH BIỂU ĐỒ USE CASE TỔNG QUÁT BẠN CHÈN VÀO ĐÂY]'
        elif 'USECASE_KHACHHANG' in line: text_ph = '[HÌNH ẢNH BIỂU ĐỒ USE CASE KHÁCH HÀNG BẠN CHÈN VÀO ĐÂY]'
        elif 'USECASE_ADMIN' in line: text_ph = '[HÌNH ẢNH BIỂU ĐỒ USE CASE ADMIN BẠN CHÈN VÀO ĐÂY]'
        elif 'ERD' in line: text_ph = '[HÌNH ẢNH BIỂU ĐỒ ERD CƠ SỞ DỮ LIỆU BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_HOME' in line: text_ph = '[HÌNH ẢNH TRANG CHỦ BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_SHOP' in line: text_ph = '[HÌNH ẢNH TRANG DANH SÁCH SẢN PHẨM BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_DETAIL' in line: text_ph = '[HÌNH ẢNH TRANG CHI TIẾT SẢN PHẨM BẠN CHÈN VÀO ĐÂY]'
        elif 'UI_ADMIN' in line: text_ph = '[HÌNH ẢNH TRANG ADMIN DASHBOARD BẠN CHÈN VÀO ĐÂY]'
        
        run_img = p_img.add_run(text_ph)
        run_img.italic = True
        run_img.bold = True
    else:
        p = doc.add_paragraph(line)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

# Apply font style
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)

output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Chuong3_PhanTichThietKe.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
