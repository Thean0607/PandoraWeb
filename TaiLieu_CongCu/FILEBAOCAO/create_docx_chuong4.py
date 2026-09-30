import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()
doc.add_heading('CHƯƠNG 4: TRIỂN KHAI HỆ THỐNG VÀ ĐÁNH GIÁ KẾT QUẢ', level=1)

text_content = """
4.1. Môi trường triển khai
4.1.1. Yêu cầu phần cứng
- Máy chủ (Server / Máy Dev): Cấu hình tối thiểu CPU Core i5, RAM 8GB, Ổ cứng SSD 256GB, kết nối mạng Internet ổn định để đảm bảo môi trường lập trình và biên dịch hệ thống.
- Máy khách (Client): Mọi thiết bị cá nhân có khả năng kết nối Internet và duyệt web (Máy tính bàn, Laptop, Máy tính bảng, Điện thoại di động).

4.1.2. Môi trường phần mềm và Công cụ phát triển
- Môi trường lập trình (IDE): Microsoft Visual Studio 2022.
- Hệ quản trị CSDL: Microsoft SQL Server Management Studio.
- Nền tảng cốt lõi: .NET Framework 4.7.2, ASP.NET MVC 5, Entity Framework 6.
- Trình duyệt kiểm thử: Google Chrome, Microsoft Edge, Safari.

4.2. Kết quả đạt được
4.2.1. Về mặt giao diện (Front-End)
Dự án đã thiết kế thành công một giao diện thương mại điện tử đồng bộ, hiện đại, mang đậm dấu ấn của thương hiệu trang sức cao cấp. Việc ứng dụng framework Bootstrap 5 giúp website tương thích hoàn hảo (Responsive Web Design) trên đa nền tảng. Các khối nội dung, hình ảnh sản phẩm và thanh điều hướng tự động co giãn, sắp xếp hợp lý trên màn hình điện thoại mà không làm vỡ bố cục, mang lại trải nghiệm xem sản phẩm xuyên suốt cho khách hàng.

4.2.2. Về mặt chức năng (Back-End)
- Phân hệ khách hàng: Hệ thống xử lý trơn tru quy trình mua sắm cốt lõi từ bước tìm kiếm sản phẩm, lọc theo chất liệu, thêm vào giỏ hàng, cho đến xác nhận thanh toán. Các tính năng bảo vệ tài khoản cá nhân, xem lịch sử đơn hàng cũng được đảm bảo tính bảo mật.
- Phân hệ quản trị: Xây dựng được trang Admin Dashboard trực quan, cho phép ban quản trị nắm bắt tình hình doanh thu tức thời. Quy trình quản lý sản phẩm, quản lý biến thể (Size, Giá) và cập nhật tiến trình vận chuyển đơn hàng được số hóa và tự động hóa đáng kể.
- Xử lý tài nguyên đám mây: Triển khai tích hợp thành công Cloudinary API, giúp việc tải lên, tối ưu dung lượng và phân phối hình ảnh diễn ra hoàn toàn tự động, giải phóng dung lượng lưu trữ cục bộ cho máy chủ.

4.3. Đánh giá hệ thống
4.3.1. Ưu điểm nổi bật
- Kiến trúc bền vững: Việc áp dụng chuẩn mô hình MVC giúp phân tách rạch ròi giữa Dữ liệu, Giao diện và Luồng điều khiển. Điều này giúp mã nguồn dự án dễ đọc, dễ bảo trì và có tiềm năng mở rộng các tính năng mới trong tương lai.
- Hiệu suất tải trang: Đẩy toàn bộ hình ảnh lên mạng lưới CDN của Cloudinary giúp giảm tải áp lực cho băng thông server nội bộ, tăng tốc độ phản hồi của trang web một cách ấn tượng.
- Trải nghiệm tương tác mượt mà: Áp dụng kỹ thuật jQuery AJAX ở nhiều chức năng (Thêm giỏ hàng, Thay đổi số lượng) giúp website cập nhật dữ liệu trực tiếp mà không cần phải tải lại toàn trang (reload), tạo cảm giác sử dụng mượt mà như một phần mềm thực thụ.

4.3.2. Những điểm còn hạn chế
Bên cạnh các kết quả đạt được, đối chiếu với một hệ thống thương mại điện tử hoàn chỉnh, dự án PandoraWeb vẫn còn một số giới hạn trong khuôn khổ đồ án:
- Chưa tích hợp cổng thanh toán trực tuyến: Hiện tại hệ thống mới chỉ hỗ trợ phương thức Thanh toán khi nhận hàng (COD - Cash on Delivery). Các giao dịch qua thẻ tín dụng quốc tế hay Ví điện tử (VNPay, MoMo) chưa được xử lý thực tế.
- Thiếu hệ thống thông báo đa kênh: Khách hàng hiện tại chỉ có thể chủ động đăng nhập web để theo dõi đơn hàng. Dự án chưa có hệ thống tự động gửi Email hóa đơn hay gửi SMS xác nhận tới số điện thoại khách hàng.
- Thiếu hệ thống gợi ý sản phẩm thông minh: Website chưa có các thuật toán phân tích hành vi người dùng (AI) để đưa ra các gợi ý "Sản phẩm có thể bạn sẽ thích" một cách tự động, cá nhân hóa.

4.4. Hướng phát triển tương lai
Dựa trên những hạn chế đã được nhìn nhận, nhóm phát triển đề xuất các định hướng nâng cấp hệ thống trong giai đoạn tiếp theo:
- Mở rộng phương thức thanh toán: Đăng ký Sandbox và tiến hành tích hợp các API thanh toán phổ biến tại Việt Nam như VNPay, ZaloPay để tạo sự tiện lợi, đa dạng hóa lựa chọn cho khách hàng và tăng tỷ lệ chuyển đổi.
- Xây dựng hệ thống Notification tự động: Bổ sung tính năng gửi Email tự động bằng SMTP (ví dụ SendGrid) để gửi mã xác minh đăng ký tài khoản, gửi hóa đơn mua hàng và gửi bản tin khuyến mãi (Newsletter).
- Nâng cấp trải nghiệm khách hàng: Tích hợp công cụ Chat trực tuyến (Live Chat) để chăm sóc khách hàng 24/7. Đồng thời, nghiên cứu áp dụng thuật toán lọc cộng tác (Collaborative Filtering) nhằm xây dựng chức năng tự động gợi ý sản phẩm liên quan.
"""

for line in text_content.strip().split('\n'):
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('4.'):
        if len(line.split('.')) == 3: # 4.1. Môi trường...
            doc.add_heading(line, level=2)
        else:
            doc.add_heading(line, level=3)
    else:
        p = doc.add_paragraph(line)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

# Apply font style
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)

output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Chuong4_TrienKhaiVaDanhGia.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
