from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

doc = Document()

# Tiêu đề
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("BÁO CÁO TIẾN ĐỘ TUẦN 11\n\n")
run.font.name = 'Times New Roman'
run.font.size = Pt(14)
run.bold = True

# Thông tin SV
info = doc.add_paragraph()
run = info.add_run("Họ và tên: Nguyễn Thế An\nMSSV: 2400004657\n")
run.font.name = 'Times New Roman'
run.font.size = Pt(12)

def add_heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

def add_text(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

def add_image_note(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"\n[ 📸 CHÈN ẢNH VÀO ĐÂY: {text} ]\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x00, 0x00, 0xFF) # Màu xanh
    r.bold = True
    r.italic = True

add_heading("1. Tổng quan tiến độ tuần 11:")
add_text("Trong tuần 11, dự án PandoraWeb đã bước sang một giai đoạn cực kỳ quan trọng: 'Tái cấu trúc kiến trúc (Refactoring) và Nâng cấp Bảo mật Toàn diện'. Thay vì tập trung thêm tính năng bề nổi, tuần này em đi sâu vào việc tối ưu hóa mã nguồn, chia tách các Controller cồng kềnh, vá các lỗ hổng bảo mật nghiêm trọng (CSRF, XSS) và cải thiện tốc độ tải trang cho người dùng. Website hiện tại đã chuyển từ mức độ 'chạy được' sang mức độ 'đạt chuẩn thực tế (Production-ready)'.")

add_heading("2. Công việc đã thực hiện trong tuần 11:")
add_text("Các đầu việc trong tuần 11 được triển khai theo hướng chuẩn hóa kiến trúc và bảo mật:")

add_text("- Tái cấu trúc khu vực Quản trị (Admin Refactoring): Đập bỏ hoàn toàn AdminController nguyên khối khổng lồ và chia tách thành các module chuyên biệt: ProductsController, OrdersController, CustomersController, MarketingController và SettingsController. Việc này giúp mã nguồn dễ đọc, dễ bảo trì và tối ưu hiệu suất biên dịch.")
add_image_note("Chụp màn hình thư mục Areas/Admin/Controllers trong Visual Studio để khoe các file Controller mới được tách ra.")

add_text("- Nâng cấp Bảo mật Mật khẩu (BCrypt): Chuyển đổi thuật toán mã hóa mật khẩu từ chuẩn SHA256 cũ kỹ sang thuật toán BCrypt hiện đại (có trộn thêm Salt ngẫu nhiên) cho cả Admin và Khách hàng. Xây dựng cơ chế Fallback thông minh giúp khách hàng cũ (dùng pass SHA256) vẫn đăng nhập được bình thường mà không bị lỗi.")
add_image_note("Mở SQL Server, chụp bảng Customers đoạn cột PasswordHash hiển thị chuỗi ngoằn ngoèo '$...' để chứng minh pass đã được mã hóa BCrypt.")

add_text("- Vá lỗ hổng bảo mật Web (CSRF & XSS):\n  + Tích hợp thẻ [ValidateAntiForgeryToken] và xây dựng cơ chế tự động gửi Token ẩn qua Javascript (fetch interceptor) cho toàn bộ các Form nhập liệu trên cả trang Quản trị và Public, ngăn chặn triệt để tấn công giả mạo yêu cầu (CSRF).\n  + Tích hợp thư viện HtmlSanitizer để 'lọc nước', xóa bỏ các thẻ <script> độc hại khỏi nội dung bài viết và mô tả sản phẩm (phòng chống tấn công XSS).")
add_image_note("Chụp màn hình file ProductsController.cs dòng có chữ [ValidateAntiForgeryToken] hoặc file _Layout.cshtml đoạn có HtmlSanitizer().")

add_text("- Cải thiện UX và Hiệu năng (Performance):\n  + Áp dụng thư viện PagedList.Mvc để Phân trang (Pagination) cho toàn bộ các danh sách trong Admin (Sản phẩm, Đơn hàng, Khách hàng...), giải quyết tình trạng đơ trang khi dữ liệu quá lớn.")
add_image_note("Mở giao diện Web phần Admin -> Quản lý sản phẩm. Chụp màn hình có chứa các nút phân trang (Trang 1, 2, 3...) ở dưới cùng của bảng.")

add_text("  + Gộp toàn bộ 10 file CSS lắt nhắt ở trang ngoài thành 1 file duy nhất (Bundling) và kích hoạt Bộ nhớ đệm (Output Caching) cho Trang chủ giúp tốc độ tải trang nhanh hơn đáng kể.")
add_image_note("Mở tab Network (F12) trên trình duyệt ở trang chủ, chụp màn hình hiển thị việc web load rất ít file CSS để chứng minh tốc độ.")

add_text("- Áp dụng 'Xóa mềm' (Soft Delete): Thay đổi logic xóa Sản phẩm. Sản phẩm giờ đây chỉ bị gán trạng thái 'deleted' (ẩn đi) chứ không bị xóa khỏi cơ sở dữ liệu, giúp bảo toàn toàn vẹn dữ liệu cho các đơn hàng cũ.")
add_image_note("Chụp màn hình bảng Products trong SQL Server, khoanh đỏ dòng có cột Status là chữ deleted.")

add_heading("3. Kết quả đạt được:")
add_text("Dự án đã lột xác hoàn toàn về mặt kiến trúc phần mềm. Việc chia nhỏ Controller giúp quá trình làm việc nhóm hoặc bảo trì sau này trở nên cực kỳ dễ dàng. Đặc biệt, website đã khắc phục được các lỗ hổng bảo mật cơ bản nhất của một trang web thương mại điện tử (XSS, CSRF). Hiệu năng ở cả khu vực Admin (nhờ phân trang) và khu vực Khách hàng (nhờ Bundling và Caching) đều được tăng tốc rõ rệt. Mã nguồn gọn gàng, sạch sẽ và an toàn hơn bao giờ hết.")

add_heading("4. Khó khăn và hướng xử lý:")
add_text("- Khó khăn: Quá trình chia tách AdminController và làm sạch code đã làm gãy rất nhiều đường dẫn (URL) và logic tham chiếu cũ, gây ra hàng loạt lỗi không tìm thấy trang (404) hoặc không tìm thấy Model (Compilation Error). Hơn nữa, việc nâng cấp mật khẩu sang BCrypt có nguy cơ làm toàn bộ khách hàng cũ không thể đăng nhập được.")
add_text("- Hướng xử lý: Em đã kiên nhẫn dò tìm và sử dụng công cụ tìm kiếm trong mã nguồn để sửa lại toàn bộ định tuyến (Areas='Admin') ở tất cả các Views. Đối với sự cố mật khẩu, em đã lập trình một cơ chế 'Kiểm tra kép' (Fallback) trong AccountController và SettingsController để cho phép hệ thống tự động nhận diện và xác thực cả 2 chuẩn mật khẩu (SHA256 và BCrypt) cùng một lúc. Các lỗi khi build dự án cũng đã được rà soát và xóa bỏ triệt để.")

add_heading("5. Kết luận:")
add_text("Tuần 11 là một bước tiến mang tính chất nền tảng và cốt lõi, tập trung vào 'chất lượng mã nguồn' và 'bảo mật hệ thống' của PandoraWeb. Với kiến trúc MVC được phân chia chuẩn mực, đi kèm các chốt chặn an ninh kiên cố và hiệu năng tải trang được tối ưu hóa, dự án đã hoàn toàn sẵn sàng cho môi trường triển khai thực tế. Trong thời gian tới, em sẽ tiến hành kiểm thử toàn diện lần cuối và chuẩn bị tài liệu báo cáo nghiệm thu tổng thể dự án.")

doc.save('c:/Users/thean/Desktop/PandoraWeb/FILEBAOCAO/Baocaotiendotuan11_NguyenTheAn_2400004657_Full.docx')
