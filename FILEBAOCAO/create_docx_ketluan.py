import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()
doc.add_heading('PHẦN KẾT LUẬN', level=1)

text_content = """
1. Kết quả đạt được của đề tài
Qua quá trình nghiên cứu, phân tích và phát triển, đề tài "Xây dựng Website Thương mại Điện tử PandoraWeb" đã cơ bản hoàn thành các mục tiêu ban đầu đề ra. Hệ thống được xây dựng và triển khai thành công dựa trên nền tảng công nghệ ASP.NET MVC 5, kết hợp cùng Entity Framework 6 và hệ quản trị cơ sở dữ liệu SQL Server.

Cụ thể, dự án đã đạt được những kết quả đáng ghi nhận sau:
- Về mặt lý thuyết: Sinh viên đã tìm hiểu, nắm bắt và vận dụng thành công kiến trúc mô hình MVC, phương pháp truy xuất dữ liệu ORM (Object-Relational Mapping), cũng như kỹ thuật tích hợp API của bên thứ ba (Cloudinary API) vào một ứng dụng thực tế.
- Về mặt thực tiễn (Sản phẩm): Xây dựng được một website bán hàng hoạt động trơn tru với luồng dữ liệu minh bạch. Đối với Khách hàng, hệ thống mang lại trải nghiệm mua sắm xuyên suốt từ khâu đăng ký, tra cứu sản phẩm đến giỏ hàng và thanh toán. Đối với Ban quản trị, dự án cung cấp một trang Dashboard mạnh mẽ để quản trị danh mục, sản phẩm, và vận hành luồng đơn hàng tự động. Đặc biệt, giao diện website được thiết kế hiện đại, tương thích hoàn hảo (Responsive) trên đa thiết bị.

2. Những mặt còn hạn chế
Mặc dù đã nỗ lực hoàn thiện hệ thống, nhưng do những giới hạn khách quan về mặt thời gian thực hiện cũng như kinh nghiệm thực tiễn, đề tài vẫn còn tồn tại một số hạn chế nhất định:
- Thiếu cổng thanh toán điện tử: Hiện tại hệ thống mới chỉ hỗ trợ phương thức Thanh toán khi nhận hàng (COD). Các giao dịch qua thẻ ngân hàng, thẻ tín dụng hoặc Ví điện tử (VNPay, MoMo) chưa được xử lý thực tế.
- Thiếu hệ thống thông báo tự động: Việc tương tác với khách hàng vẫn còn bị động. Dự án chưa có hệ thống tự động gửi Email xác nhận hóa đơn hay gửi SMS thông báo khi trạng thái đơn hàng thay đổi.
- Thiếu tính năng cá nhân hóa: Website chưa được ứng dụng các thuật toán phân tích dữ liệu để gợi ý sản phẩm tự động dựa trên sở thích và hành vi của người tiêu dùng.

3. Hướng phát triển trong tương lai
Từ những khuyết điểm đã tự nhìn nhận, nhằm đưa hệ thống ngày càng hoàn thiện và tiệm cận với một sản phẩm thương mại thực thụ, nhóm phát triển đề ra các định hướng nâng cấp trong tương lai như sau:
- Tích hợp thanh toán số: Tiến hành đăng ký môi trường Sandbox và lập trình tích hợp các API thanh toán phổ biến tại Việt Nam (như VNPay, ZaloPay, PayPal) nhằm đa dạng hóa lựa chọn cho khách mua hàng.
- Xây dựng hệ thống Notification tự động: Bổ sung module gửi Email tự động bằng giao thức SMTP (ví dụ: SendGrid) để gửi mã xác thực (OTP), hóa đơn điện tử và các chiến dịch Marketing định kỳ.
- Nâng cấp Trí tuệ nhân tạo (AI): Nghiên cứu và áp dụng thuật toán lọc cộng tác (Collaborative Filtering) để xây dựng chức năng "Gợi ý sản phẩm" thông minh. Đồng thời, triển khai tích hợp Chatbot trực tuyến hỗ trợ giải đáp thắc mắc của khách hàng 24/7.
- Tối ưu hóa phân tích dữ liệu: Mở rộng các chức năng trong trang Admin Dashboard, cung cấp các biểu đồ thống kê chuyên sâu hơn (như thống kê doanh thu theo khu vực địa lý, độ tuổi khách hàng) nhằm hỗ trợ đưa ra các chiến lược kinh doanh hiệu quả.
"""

for line in text_content.strip().split('\n'):
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('1.') or line.startswith('2.') or line.startswith('3.'):
        doc.add_heading(line, level=2)
    else:
        p = doc.add_paragraph(line)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

# Apply font style
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)

output_path = r'c:\Users\thean\Desktop\PandoraWeb\FILEBAOCAO\Phan_Ket_Luan.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
